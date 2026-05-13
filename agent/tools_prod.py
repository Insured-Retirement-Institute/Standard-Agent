"""
IRI Standards Agent — Production Tool Definitions (v1.0)
Only tools the LLM literally cannot do itself (HTTP fetch).

Tools (2):
  - fetch_yaml_from_url — download YAML from a URL
  - fetch_data_dictionary — download/parse Excel/CSV Data Dictionary

Dependencies: PyYAML, strands (NO ruamel, NO openapi-spec-validator, NO jsonschema)
"""

import json
import os
import re
import tempfile
from typing import Optional

import yaml
from strands import tool
import urllib.request
import urllib.error


@tool
def fetch_yaml_from_url(url: str) -> str:
    """Fetch YAML/OpenAPI content from a URL (GitHub, raw links, or any public URL).
    Automatically converts GitHub blob URLs to raw content URLs.
    Use this when the user provides a link to a YAML file instead of pasting content.

    Args:
        url: URL to the YAML file (GitHub blob, raw, or any public URL)
    """
    try:
        original_url = url.strip()

        # Convert GitHub blob URLs to raw content URLs
        raw_url = original_url
        github_blob_match = re.match(
            r"https?://github\.com/([^/]+)/([^/]+)/blob/(.+)", original_url
        )
        if github_blob_match:
            org, repo, rest = github_blob_match.groups()
            raw_url = f"https://raw.githubusercontent.com/{org}/{repo}/{rest}"

        github_tree_match = re.match(
            r"https?://github\.com/([^/]+)/([^/]+)/tree/(.+)", original_url
        )
        if github_tree_match:
            return json.dumps({
                "success": False,
                "error": "This appears to be a directory URL, not a file. Please provide a direct link to a .yaml or .yml file."
            })

        req = urllib.request.Request(raw_url, headers={"User-Agent": "IRI-Standards-Agent/1.0"})
        with urllib.request.urlopen(req, timeout=30) as resp:
            charset = resp.headers.get_content_charset() or "utf-8"
            yaml_content = resp.read().decode(charset)

        # Validate it's parseable YAML
        try:
            parsed = yaml.safe_load(yaml_content)
            if not isinstance(parsed, dict):
                return json.dumps({
                    "success": False,
                    "error": "URL content parsed but is not a YAML dictionary/mapping."
                })
            is_openapi = "openapi" in parsed
        except yaml.YAMLError as ye:
            return json.dumps({
                "success": False,
                "error": f"URL content is not valid YAML: {str(ye)[:300]}"
            })

        info = parsed.get("info", {}) if is_openapi else {}
        return json.dumps({
            "success": True,
            "source_url": original_url,
            "resolved_url": raw_url,
            "is_openapi": is_openapi,
            "openapi_version": parsed.get("openapi", "N/A") if is_openapi else "N/A",
            "title": info.get("title", "N/A"),
            "version": info.get("version", "N/A"),
            "line_count": yaml_content.count("\n") + 1,
            "char_count": len(yaml_content),
            "yaml_content": yaml_content
        })

    except urllib.error.HTTPError as e:
        return json.dumps({"success": False, "error": f"HTTP {e.code}: {e.reason}."})
    except urllib.error.URLError as e:
        return json.dumps({"success": False, "error": f"URL error: {str(e.reason)}."})
    except Exception as e:
        return json.dumps({"success": False, "error": f"Failed to fetch: {str(e)[:300]}"})


@tool
def fetch_data_dictionary(url: str, sheet_name: str = "", field_name_column: str = "") -> str:
    """Fetch a Data Dictionary Excel (.xlsx) or CSV file from a URL and return
    its contents as a JSON array of field definitions.

    Args:
        url: URL to the Data Dictionary file (GitHub blob, raw, or any public URL)
        sheet_name: Optional Excel sheet name. If empty, uses the first sheet.
        field_name_column: Optional column name containing field names. Auto-detects if empty.
    """
    try:
        original_url = url.strip()

        raw_url = original_url
        github_blob_match = re.match(
            r"https?://github\.com/([^/]+)/([^/]+)/blob/(.+)", original_url
        )
        if github_blob_match:
            org, repo, rest = github_blob_match.groups()
            raw_url = f"https://raw.githubusercontent.com/{org}/{repo}/{rest}"

        req = urllib.request.Request(raw_url, headers={"User-Agent": "IRI-Standards-Agent/1.0"})
        with urllib.request.urlopen(req, timeout=60) as resp:
            file_bytes = resp.read()

        lower_url = raw_url.lower()
        is_csv = lower_url.endswith(".csv")
        is_excel = any(lower_url.endswith(ext) for ext in (".xlsx", ".xls", ".xlsm"))

        if not is_csv and not is_excel:
            try:
                file_bytes.decode("utf-8")
                is_csv = True
            except UnicodeDecodeError:
                is_excel = True

        rows = []

        if is_excel:
            try:
                import openpyxl
            except ImportError:
                return json.dumps({"success": False, "error": "openpyxl not installed. Run: pip install openpyxl"})

            with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
                tmp.write(file_bytes)
                tmp_path = tmp.name

            try:
                wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
                if sheet_name:
                    if sheet_name not in wb.sheetnames:
                        return json.dumps({"success": False, "error": f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}"})
                    ws = wb[sheet_name]
                else:
                    ws = wb.active

                all_rows = []
                for row in ws.iter_rows(values_only=True):
                    all_rows.append([str(c).strip() if c is not None else "" for c in row])
                wb.close()
            finally:
                os.unlink(tmp_path)

            if not all_rows:
                return json.dumps({"success": False, "error": "Sheet is empty"})

            headers = all_rows[0]
            for i, row in enumerate(all_rows):
                if any(cell for cell in row):
                    headers = row
                    all_rows = all_rows[i+1:]
                    break

            headers = [h.strip().lower().replace(" ", "_").replace("-", "_") for h in headers]
            for row in all_rows:
                if not any(cell for cell in row):
                    continue
                record = {}
                for j, val in enumerate(row):
                    if j < len(headers) and headers[j]:
                        record[headers[j]] = val
                rows.append(record)

        elif is_csv:
            import csv
            import io
            text = file_bytes.decode("utf-8-sig")
            reader = csv.reader(io.StringIO(text))
            all_rows = list(reader)

            if not all_rows:
                return json.dumps({"success": False, "error": "CSV is empty"})

            headers = [h.strip().lower().replace(" ", "_").replace("-", "_") for h in all_rows[0]]
            for row in all_rows[1:]:
                if not any(cell.strip() for cell in row):
                    continue
                record = {}
                for j, val in enumerate(row):
                    if j < len(headers) and headers[j]:
                        record[headers[j]] = val.strip()
                rows.append(record)

        # Auto-detect field_name column
        fn_col = field_name_column.strip().lower().replace(" ", "_").replace("-", "_") if field_name_column else ""
        if not fn_col and rows:
            candidates = [k for k in rows[0].keys()
                         if any(term in k for term in ["field", "name", "property", "element", "attribute", "column"])]
            if candidates:
                fn_col = candidates[0]

        if fn_col and rows and fn_col in rows[0]:
            for row in rows:
                row["field_name"] = row.get(fn_col, "")

        return json.dumps({
            "success": True,
            "source_url": original_url,
            "format": "excel" if is_excel else "csv",
            "row_count": len(rows),
            "columns": list(rows[0].keys()) if rows else [],
            "field_name_column": fn_col or "(not detected)",
            "sample_rows": rows[:3],
            "data": rows
        })

    except urllib.error.HTTPError as e:
        return json.dumps({"success": False, "error": f"HTTP {e.code}: {e.reason}"})
    except Exception as e:
        return json.dumps({"success": False, "error": f"Failed: {str(e)[:300]}"})


# ============================================================================
# TOOL REGISTRY
# ============================================================================
PROD_TOOLS = [
    fetch_yaml_from_url,
    fetch_data_dictionary,
]
